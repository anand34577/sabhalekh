"""Transcript exports, UTF-8 everywhere: TXT, Markdown, SRT, VTT, JSON, CSV, DOCX and PDF.

PDF: fpdf2 with HarfBuzz text shaping (uharfbuzz) and the bundled Noto Sans / Noto Sans Devanagari fonts, so
Devanagari conjuncts and matras are shaped correctly (a naive PDF library would print unjoined code points).
DOCX: python-docx with Noto Sans for Latin and Noto Sans Devanagari as the complex-script font (w:cs), language
tagged hi-IN, so Word picks the right font for Hindi runs.
"""
from __future__ import annotations

import csv
import io
import json
import time
import zipfile
from dataclasses import dataclass

from . import config
from . import text as textutil

FORMATS = {"txt": "text/plain", "md": "text/markdown", "srt": "application/x-subrip", "vtt": "text/vtt",
           "json": "application/json", "csv": "text/csv",
           "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
           "pdf": "application/pdf"}


@dataclass
class ExportOptions:
    timestamps: bool = True
    speakers: bool = True
    merge: bool = True  # merge consecutive lines of the same speaker


def ts(sec: float, millis: bool = False, sep: str = ",") -> str:
    sec = max(0.0, sec)
    h, rem = divmod(int(sec), 3600)
    m, s = divmod(rem, 60)
    if millis:
        return f"{h:02d}:{m:02d}:{s:02d}{sep}{int(round((sec - int(sec)) * 1000)) % 1000:03d}"
    return f"{h:02d}:{m:02d}:{s:02d}"


def build_lines(doc: dict, opts: ExportOptions) -> list[dict]:
    """[{start, end, speaker, color, text, language}] in time order, merged per speaker if requested."""
    spk = {s["id"]: s for s in doc["speakers"]}
    lines: list[dict] = []
    for seg in doc["segments"]:
        sp = spk.get(seg["speaker_id"])
        text = textutil.nfc(seg["display_text"]).strip()
        if not text:
            continue
        line = {"start": seg["start_sec"], "end": seg["end_sec"], "speaker": sp["display_name"] if sp else "Unknown",
                "color": sp["color"] if sp else "#888888", "text": text, "language": seg.get("language"),
                "overlap": bool(seg.get("overlap_flag"))}
        if opts.merge and lines and lines[-1]["speaker"] == line["speaker"] and line["start"] - lines[-1]["end"] < 5:
            lines[-1]["text"] += " " + text
            lines[-1]["end"] = line["end"]
        else:
            lines.append(line)
    return lines


def _prefix(line: dict, opts: ExportOptions) -> str:
    p = ""
    if opts.timestamps:
        p += f"[{ts(line['start'])}] "
    if opts.speakers:
        p += f"{line['speaker']}: "
    return p


def to_txt(doc: dict, opts: ExportOptions) -> bytes:
    lines = [f"{_prefix(l, opts)}{l['text']}" for l in build_lines(doc, opts)]
    return ("\n".join(lines) + "\n").encode("utf-8")


def to_md(doc: dict, opts: ExportOptions) -> bytes:
    s = doc["session"]
    out = [f"# {s['title']}", "", f"*{time.strftime('%d %B %Y, %H:%M', time.localtime(s['created_at']))} · "
                                  f"{ts(s['duration_sec'] or 0)}*", ""]
    if doc.get("talk_time"):
        out += ["**Talk time:** " + " · ".join(f"{t['name']} {t['percent']:.0f}%" for t in doc["talk_time"]), ""]
    for l in build_lines(doc, opts):
        head = []
        if opts.timestamps:
            head.append(f"`{ts(l['start'])}`")
        if opts.speakers:
            head.append(f"**{l['speaker']}**")
        out.append((" ".join(head) + " " if head else "") + l["text"])
        out.append("")
    return "\n".join(out).encode("utf-8")


def to_srt(doc: dict, opts: ExportOptions) -> bytes:
    out = []
    for i, l in enumerate(build_lines(doc, ExportOptions(opts.timestamps, opts.speakers, False)), 1):
        who = f"{l['speaker']}: " if opts.speakers else ""
        out.append(f"{i}\n{ts(l['start'], True)} --> {ts(l['end'], True)}\n{who}{l['text']}\n")
    return "\n".join(out).encode("utf-8")


def to_vtt(doc: dict, opts: ExportOptions) -> bytes:
    out = ["WEBVTT", ""]
    for l in build_lines(doc, ExportOptions(opts.timestamps, opts.speakers, False)):
        who = f"<v {l['speaker']}>" if opts.speakers else ""
        out.append(f"{ts(l['start'], True, '.')} --> {ts(l['end'], True, '.')}\n{who}{l['text']}\n")
    return "\n".join(out).encode("utf-8")


def to_json(doc: dict, opts: ExportOptions) -> bytes:
    s = doc["session"]
    spk = {x["id"]: x["display_name"] for x in doc["speakers"]}
    data = {
        "title": s["title"], "created_at": s["created_at"], "duration_sec": s["duration_sec"],
        "source_type": s["source_type"], "language_mode": s["language_mode"],
        "detected_languages": s.get("detected_languages", []), "asr_model": s.get("asr_model"),
        "diarization_engine": s.get("diarization_engine"),
        "speakers": [{"name": x["display_name"], "label": x["label"], "color": x["color"]} for x in doc["speakers"]],
        "talk_time": doc.get("talk_time", []),
        "segments": [{"start": g["start_sec"], "end": g["end_sec"],
                      "speaker": spk.get(g["speaker_id"]) if opts.speakers else None,
                      "text": g["display_text"], "language": g["language"], "confidence": g["confidence"],
                      "overlap": bool(g["overlap_flag"]), "edited": g["edited_text"] is not None,
                      "words": g["words"]} for g in doc["segments"]],
    }
    return json.dumps(data, ensure_ascii=False, indent=2).encode("utf-8")


def to_csv(doc: dict, opts: ExportOptions) -> bytes:
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow(["start", "end", "speaker", "text", "language"])
    for l in build_lines(doc, opts):
        w.writerow([ts(l["start"], True) if opts.timestamps else "", ts(l["end"], True) if opts.timestamps else "",
                    l["speaker"] if opts.speakers else "", l["text"], l["language"] or ""])
    return ("﻿" + buf.getvalue()).encode("utf-8")  # BOM so Excel opens UTF-8 Hindi correctly


# ------------------------------------------------------------------ DOCX
def to_docx(doc: dict, opts: ExportOptions) -> bytes:
    from docx import Document
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Pt, RGBColor

    def fonts(run, size=None, bold=None, color=None):
        run.font.name = "Noto Sans"
        rpr = run._element.get_or_add_rPr()
        rfonts = rpr.find(qn("w:rFonts"))
        if rfonts is None:
            rfonts = OxmlElement("w:rFonts")
            rpr.insert(0, rfonts)
        for attr in ("w:ascii", "w:hAnsi", "w:eastAsia"):
            rfonts.set(qn(attr), "Noto Sans")
        rfonts.set(qn("w:cs"), "Noto Sans Devanagari")
        lang = OxmlElement("w:lang")
        lang.set(qn("w:val"), "en-IN")
        lang.set(qn("w:bidi"), "hi-IN")
        rpr.append(lang)
        if size:
            run.font.size = Pt(size)
            szcs = OxmlElement("w:szCs")
            szcs.set(qn("w:val"), str(int(size * 2)))
            rpr.append(szcs)
        if bold:
            run.bold = True
            rpr.append(OxmlElement("w:bCs"))
        if color:
            run.font.color.rgb = RGBColor.from_string(color.lstrip("#").upper())

    s = doc["session"]
    d = Document()
    st = d.styles["Normal"]
    st.font.name = "Noto Sans"
    st.font.size = Pt(10.5)
    st.element.rPr.rFonts.set(qn("w:cs"), "Noto Sans Devanagari")
    fonts(d.add_paragraph().add_run(s["title"]), size=20, bold=True)
    meta = (f"{time.strftime('%d %B %Y, %H:%M', time.localtime(s['created_at']))}  ·  {ts(s['duration_sec'] or 0)}"
            f"  ·  {'Recorded' if s['source_type'] == 'recorded' else 'Imported'}")
    fonts(d.add_paragraph().add_run(meta), size=9, color="#666666")
    if doc.get("talk_time"):
        fonts(d.add_paragraph().add_run("Speakers"), size=13, bold=True)
        for t in doc["talk_time"]:
            p = d.add_paragraph()
            fonts(p.add_run("● "), color=t["color"])
            fonts(p.add_run(f"{t['name']}  "), bold=True)
            fonts(p.add_run(f"{t['percent']:.0f}% · {ts(t['seconds'])}"), color="#666666")
        fonts(d.add_paragraph().add_run("Transcript"), size=13, bold=True)
    for l in build_lines(doc, opts):
        p = d.add_paragraph()
        p.paragraph_format.space_after = Pt(6)
        if opts.timestamps:
            fonts(p.add_run(f"[{ts(l['start'])}] "), size=9, color="#888888")
        if opts.speakers:
            fonts(p.add_run(f"{l['speaker']}: "), bold=True, color=l["color"])
        fonts(p.add_run(l["text"]))
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


# ------------------------------------------------------------------ PDF
def _pdf():
    from fpdf import FPDF

    f = config.FONTS_DIR
    pdf = FPDF(format="A4", unit="mm")
    pdf.set_auto_page_break(True, margin=18)
    pdf.add_font("NotoSans", "", str(f / "NotoSans-Regular.ttf"))
    pdf.add_font("NotoSans", "B", str(f / "NotoSans-Bold.ttf"))
    pdf.add_font("NotoDeva", "", str(f / "NotoSansDevanagari-Regular.ttf"))
    pdf.add_font("NotoDeva", "B", str(f / "NotoSansDevanagari-Bold.ttf"))
    pdf.set_fallback_fonts(["NotoDeva"], exact_match=False)
    pdf.set_text_shaping(True)
    return pdf


def _rgb(hex_color: str) -> tuple[int, int, int]:
    h = hex_color.lstrip("#")
    return int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16)


def to_pdf(doc: dict, opts: ExportOptions) -> bytes:
    s = doc["session"]
    pdf = _pdf()
    pdf.set_title(s["title"])
    pdf.set_creator("Sabhalekh (offline)")
    pdf.add_page()
    pdf.set_margins(18, 18, 18)
    pdf.set_font("NotoSans", "B", 20)
    pdf.multi_cell(0, 10, s["title"], new_x="LMARGIN", new_y="NEXT")
    pdf.set_font("NotoSans", "", 9)
    pdf.set_text_color(110, 110, 110)
    meta = (f"{time.strftime('%d %B %Y, %H:%M', time.localtime(s['created_at']))}  ·  {ts(s['duration_sec'] or 0)}"
            f"  ·  {'Recorded' if s['source_type'] == 'recorded' else 'Imported'}"
            f"  ·  {', '.join(x.upper() for x in s.get('detected_languages') or []) or s['language_mode']}")
    pdf.cell(0, 6, meta, new_x="LMARGIN", new_y="NEXT")
    pdf.ln(4)
    tt = doc.get("talk_time") or []
    if tt:
        pdf.set_text_color(20, 20, 20)
        pdf.set_font("NotoSans", "B", 12)
        pdf.cell(0, 8, "Speaker summary", new_x="LMARGIN", new_y="NEXT")
        w = pdf.w - pdf.l_margin - pdf.r_margin
        x0, y0 = pdf.get_x(), pdf.get_y()
        for t in tt:  # stacked talk-time bar
            seg_w = w * t["percent"] / 100
            if seg_w > 0:
                pdf.set_fill_color(*_rgb(t["color"]))
                pdf.rect(x0, y0, seg_w, 4, style="F")
                x0 += seg_w
        pdf.ln(7)
        pdf.set_font("NotoSans", "", 10)
        for t in tt:
            pdf.set_fill_color(*_rgb(t["color"]))
            pdf.rect(pdf.get_x(), pdf.get_y() + 1.6, 3, 3, style="F")
            pdf.set_x(pdf.get_x() + 5)
            pdf.cell(70, 6, t["name"])
            pdf.set_text_color(110, 110, 110)
            pdf.cell(0, 6, f"{t['percent']:.0f}%   {ts(t['seconds'])}", new_x="LMARGIN", new_y="NEXT")
            pdf.set_text_color(20, 20, 20)
        pdf.ln(4)
        pdf.set_font("NotoSans", "B", 12)
        pdf.cell(0, 8, "Transcript", new_x="LMARGIN", new_y="NEXT")
    for l in build_lines(doc, opts):
        pdf.set_font("NotoSans", "", 8.5)
        if opts.timestamps or opts.speakers:
            if opts.timestamps:
                pdf.set_text_color(130, 130, 130)
                pdf.cell(18, 5.5, ts(l["start"]))
            if opts.speakers:
                pdf.set_font("NotoSans", "B", 9.5)
                pdf.set_text_color(*_rgb(l["color"]))
                pdf.cell(0, 5.5, l["speaker"], new_x="LMARGIN", new_y="NEXT")
            else:
                pdf.ln(5.5)
        pdf.set_font("NotoSans", "", 10.5)
        pdf.set_text_color(25, 25, 25)
        pdf.multi_cell(0, 6.2, l["text"], new_x="LMARGIN", new_y="NEXT")
        pdf.ln(2.2)
    pdf.set_y(-14)
    return bytes(pdf.output())


EXPORTERS = {"txt": to_txt, "md": to_md, "srt": to_srt, "vtt": to_vtt, "json": to_json, "csv": to_csv,
             "docx": to_docx, "pdf": to_pdf}


def export(doc: dict, fmt: str, opts: ExportOptions) -> bytes:
    if fmt not in EXPORTERS:
        raise ValueError(f"unknown export format {fmt!r}")
    return EXPORTERS[fmt](doc, opts)


def safe_filename(title: str, ext: str) -> str:
    import unicodedata

    def ok(c: str) -> bool:  # letters, combining marks (Devanagari matras/virama), digits
        return unicodedata.category(c)[0] in "LMN" or c in " -_.()"

    keep = "".join(c if ok(c) else "_" for c in textutil.nfc(title)).strip() or "transcript"
    return f"{keep[:80]}.{ext}"


def export_zip(docs: list[dict], formats: list[str], opts: ExportOptions) -> bytes:
    buf = io.BytesIO()
    used: set[str] = set()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
        for doc in docs:
            for fmt in formats:
                name = safe_filename(doc["session"]["title"], fmt)
                base, n = name, 2
                while name in used:
                    name = base.replace(f".{fmt}", f" ({n}).{fmt}")
                    n += 1
                used.add(name)
                z.writestr(name, export(doc, fmt, opts))
    return buf.getvalue()
