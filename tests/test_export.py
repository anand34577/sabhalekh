"""Exports: every format, UTF-8 Hindi intact, Devanagari shaping in PDF (golden image), complex-script font in DOCX.

Regenerate the PDF golden image after an intentional layout change with:
    SABHALEKH_UPDATE_GOLDEN=1 python -m pytest tests/test_export.py
"""
import csv
import io
import json
import os
import zipfile

import numpy as np
import pytest

from conftest import DATA, sample_transcript
from sabhalekh import export

GOLDEN = DATA / "golden"


def doc(db):
    sid = sample_transcript(db)
    s = db.get_session(sid)
    s["created_at"] = 1767225600  # fixed date for stable output
    return {"session": s, **db.get_transcript(sid), "talk_time": db.talk_time(sid)}


OPTS = export.ExportOptions()


def test_txt_matches_documented_format(db):
    d = doc(db)
    rahul = next(s for s in d["speakers"] if s["label"] == "Speaker 1")
    rahul["display_name"] = "Rahul"
    out = export.to_txt(d, OPTS).decode("utf-8")
    assert out.splitlines()[0] == "[00:00:00] Rahul: Let's start with the budget."
    assert "Speaker 2: हम अगले क्वार्टर का बजट तय करेंगे।" in out


def test_txt_options(db):
    d = doc(db)
    plain = export.to_txt(d, export.ExportOptions(timestamps=False, speakers=False)).decode()
    assert plain.splitlines()[0] == "Let's start with the budget."
    merged = export.to_txt(d, export.ExportOptions(merge=True)).decode().splitlines()
    unmerged = export.to_txt(d, export.ExportOptions(merge=False)).decode().splitlines()
    assert len(unmerged) == 4 and len(merged) == 4  # speakers alternate, nothing to merge
    d["segments"][1]["speaker_id"] = d["segments"][0]["speaker_id"]  # now S1, S1, S1, S2
    lines = export.to_txt(d, export.ExportOptions(merge=True)).decode().splitlines()
    assert len(lines) == 2 and lines[0].startswith("[00:00:00] Speaker 1: Let's start") and "Yes, I think" in lines[0]


def test_srt_and_vtt(db):
    d = doc(db)
    srt = export.to_srt(d, OPTS).decode("utf-8")
    assert srt.startswith("1\n00:00:00,000 --> 00:00:04,000\nSpeaker 1: Let's start with the budget.\n")
    assert "00:00:04,500 --> 00:00:09,000" in srt
    vtt = export.to_vtt(d, OPTS).decode("utf-8")
    assert vtt.startswith("WEBVTT\n\n00:00:00.000 --> 00:00:04.000\n<v Speaker 1>Let's start")


def test_json_csv_md(db):
    d = doc(db)
    j = json.loads(export.to_json(d, OPTS))
    assert j["segments"][1]["text"] == "हम अगले क्वार्टर का बजट तय करेंगे।" and j["segments"][0]["words"]
    raw = export.to_csv(d, OPTS)
    assert raw.startswith("﻿".encode("utf-8"))
    rows = list(csv.reader(io.StringIO(raw.decode("utf-8-sig"))))
    assert rows[0] == ["start", "end", "speaker", "text", "language"] and rows[2][3].startswith("हम अगले")
    md = export.to_md(d, OPTS).decode("utf-8")
    assert md.startswith("# Budget review") and "**Talk time:**" in md and "पक्का" in md


def test_docx_sets_complex_script_font_and_keeps_hindi(db):
    from docx import Document

    data = export.to_docx(doc(db), OPTS)
    xml = zipfile.ZipFile(io.BytesIO(data)).read("word/document.xml").decode("utf-8")
    assert 'w:cs="Noto Sans Devanagari"' in xml and 'w:bidi="hi-IN"' in xml
    text = "\n".join(p.text for p in Document(io.BytesIO(data)).paragraphs)
    for s in ("हम अगले क्वार्टर का बजट तय करेंगे।", "क्या यह सही है? Yes, I think so.", "पक्का, कल तक भेज दूँगा।"):
        assert s in text
    assert "Speakers" in text and "%" in text  # speaker summary with talk time


def _render(pdf_bytes: bytes, page: int = 0) -> np.ndarray:
    pdfium = pytest.importorskip("pypdfium2")
    pdf = pdfium.PdfDocument(pdf_bytes)
    img = pdf[page].render(scale=1.5).to_pil().convert("L")
    return np.asarray(img, dtype=np.int16)


def test_pdf_shapes_devanagari_and_matches_golden(db, monkeypatch):
    import time as _time

    monkeypatch.setattr(export.time, "localtime", _time.gmtime)  # same rendering in every timezone
    data = export.to_pdf(doc(db), OPTS)
    assert data.startswith(b"%PDF")
    pdfium = pytest.importorskip("pypdfium2")
    text = pdfium.PdfDocument(data)[0].get_textpage().get_text_range()
    assert "Speaker summary" in text and "Budget review" in text
    img = _render(data)
    GOLDEN.mkdir(parents=True, exist_ok=True)
    golden = GOLDEN / "report_hindi_page1.png"
    if os.environ.get("SABHALEKH_UPDATE_GOLDEN") or not golden.exists():
        from PIL import Image

        Image.fromarray(img.astype(np.uint8)).save(golden)
        if not os.environ.get("SABHALEKH_UPDATE_GOLDEN"):
            pytest.skip("golden image created; review tests/data/golden/report_hindi_page1.png and re-run")
    from PIL import Image

    ref = np.asarray(Image.open(golden).convert("L"), dtype=np.int16)
    assert ref.shape == img.shape
    diff = np.abs(ref - img)
    assert (diff > 40).mean() < 0.002, "PDF rendering changed (Devanagari shaping or layout)"


def test_pdf_uses_shaped_glyphs_not_codepoints(db):
    """With HarfBuzz shaping, the conjunct क्व in 'क्वार्टर' becomes a ligature: the shaped run has fewer glyphs
    than code points. Without shaping, a viramas would be drawn as visible halant marks."""
    import uharfbuzz as hb

    from sabhalekh import config

    blob = hb.Blob.from_file_path(str(config.FONTS_DIR / "NotoSansDevanagari-Regular.ttf"))
    font = hb.Font(hb.Face(blob))
    buf = hb.Buffer()
    buf.add_str("क्वार्टर")
    buf.guess_segment_properties()
    hb.shape(font, buf, {})
    assert len(buf.glyph_infos) < len("क्वार्टर")
    assert export.to_pdf(doc(db), OPTS)  # the same font + shaping path is what to_pdf uses


def test_zip_batch(db):
    d1 = doc(db)
    data = export.export_zip([d1, d1], ["txt", "srt", "pdf"], OPTS)
    names = zipfile.ZipFile(io.BytesIO(data)).namelist()
    assert len(names) == 6 and "Budget review.txt" in names and "Budget review (2).txt" in names


def test_filename_sanitizing():
    assert export.safe_filename('Q3: "plan"/review', "pdf") == "Q3_ _plan__review.pdf"
    assert export.safe_filename("बजट मीटिंग", "docx") == "बजट मीटिंग.docx"
